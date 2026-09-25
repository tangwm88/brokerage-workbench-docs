#!/usr/bin/env python3
"""MD转Word"""
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import re
import urllib.request

doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = 'Microsoft YaHei'
font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

import sys
md_file = sys.argv[1] if len(sys.argv) > 1 else '资产管理自动驾驶平台设计方案 v3.0.md'
docx_file = sys.argv[2] if len(sys.argv) > 2 else md_file.replace('.md', '.docx')

with open(md_file, 'r', encoding='utf-8') as f:
    lines = f.readlines()

i = 0
while i < len(lines):
    line = lines[i].rstrip()
    
    if not line:
        i += 1
        continue
    
    if line.startswith('# '):
        doc.add_heading(line[2:], level=1)
        i += 1
        continue
    elif line.startswith('## '):
        doc.add_heading(line[3:], level=2)
        i += 1
        continue
    elif line.startswith('### '):
        doc.add_heading(line[4:], level=3)
        i += 1
        continue
    elif line.startswith('#### '):
        doc.add_heading(line[5:], level=4)
        i += 1
        continue
    
    if line.startswith('> '):
        p = doc.add_paragraph(line[2:])
        p.paragraph_format.left_indent = Inches(0.3)
        i += 1
        continue
    if line.startswith('>'):
        p = doc.add_paragraph(line[1:].strip())
        p.paragraph_format.left_indent = Inches(0.3)
        i += 1
        continue
    
    img_match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
    if img_match:
        alt, url = img_match.groups()
        if 'github' in url:
            try:
                urllib.request.urlretrieve(url, '/tmp/diagram.png')
                doc.add_picture('/tmp/diagram.png', width=Inches(6))
                doc.add_paragraph()
            except:
                doc.add_paragraph(f'[图片: {alt}]')
        else:
            doc.add_paragraph(f'[图片: {alt}]')
        i += 1
        continue
    
    if line.startswith('|') and i + 1 < len(lines) and lines[i+1].strip().startswith('|'):
        table_lines = []
        while i < len(lines) and lines[i].strip().startswith('|'):
            table_lines.append(lines[i].strip())
            i += 1
        
        rows = []
        for tl in table_lines:
            if '---' in tl:
                continue
            cells = [c.strip() for c in tl.split('|')[1:-1]]
            if cells:
                rows.append(cells)
        
        if rows:
            num_cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=num_cols)
            table.style = 'Table Grid'
            for ri, row in enumerate(rows):
                for ci in range(num_cols):
                    if ci < len(row):
                        table.rows[ri].cells[ci].text = row[ci]
            doc.add_paragraph()
        continue
    
    if line.strip().startswith('```'):
        code_lines = []
        i += 1
        while i < len(lines) and not lines[i].strip().startswith('```'):
            code_lines.append(lines[i].rstrip())
            i += 1
        i += 1
        for cl in code_lines:
            p = doc.add_paragraph(cl)
            for run in p.runs:
                run.font.name = 'Courier New'
                run.font.size = Pt(9)
        continue
    
    p = doc.add_paragraph(line)
    i += 1

output_path = docx_file
doc.save(output_path)
print(f'OK: {output_path}')
